"""
api/excel.py  —  Vercel Serverless Function (Python)
Genera Star Performance o Matriz en .xlsx con colores completos.
Recibe POST JSON: { tipo: "star" | "matriz", sucursal, entrenador, empleados: [...] }
Devuelve: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet
"""
from http.server import BaseHTTPRequestHandler
import json, base64, io, datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── PALETA ────────────────────────────────────────────────────
NEG   = "0D0D0D"; NEG2  = "181818"; NEG3  = "222222"
AMAR  = "FFDE21"; ROJO  = "C8102E"; AZUL  = "1565C0"
VERDE = "2E7D32"; BL8B  = "8B0000"
BLANC = "F5F5F5"; GRIS  = "888888"; GRIS2 = "444444"

# ── COLUMNAS ─────────────────────────────────────────────────
SP_GRUPOS = [
    {"id":"ob","label":"ORIENTACIÓN",             "color":ROJO, "sub":"","cols":[
        {"id":"ob_onboarding","label":"Onboarding /\nOrientación"},
    ]},
    {"id":"bl","label":"BACKLINE / PRODUCCIÓN",   "color":BL8B, "sub":"","cols":[
        {"id":"bl_freidoras","label":"Fry\nStation"},
        {"id":"bl_feeder",   "label":"Feeder\nStation"},
        {"id":"bl_cocina",   "label":"Cook\nStation"},
        {"id":"bl_tenders",  "label":"Chicken Tenders\nStation"},
    ]},
    {"id":"fl","label":"FRONTLINE / HOSPITALIDAD","color":AZUL, "sub":"","cols":[
        {"id":"fl_comedor",      "label":"Dining Room\nStation"},
        {"id":"fl_cajero",       "label":"Cashier\nStation"},
        {"id":"fl_autoservicio", "label":"Drive-Thru\nStation"},
        {"id":"fl_drive_speed",  "label":"Drive-Thru Speed\nCertification"},
    ]},
    {"id":"ct","label":"CREW TRAINER",            "color":VERDE,"sub":"","cols":[
        {"id":"ct_backline", "label":"Crew Trainer –\nBACKLINE"},
        {"id":"ct_frontline","label":"Crew Trainer –\nFRONTLINE"},
        {"id":"ct_pic",      "label":"Centerpost\n(PIC)"},
    ]},
]

MZ_GRUPOS = [
    {"id":"ob","label":"ORIENTACIÓN",  "color":ROJO, "sub":"","cols":[
        {"id":"ob_onboarding","label":"Onboarding / Orientación\nChecklist Completed"},
    ]},
    {"id":"bl","label":"BACKLINE",      "color":BL8B, "sub":"PRODUCTION","cols":[
        {"id":"bl_freidoras",  "label":"Fry Station"},
        {"id":"bl_feeder",     "label":"Feeder Station"},
        {"id":"bl_cocina",     "label":"Cook Station"},
        {"id":"bl_grill",      "label":"Grill Station"},
        {"id":"bl_burrito",    "label":"Green Burrito Station"},
        {"id":"bl_tenders",    "label":"Chicken Tenders Station"},
        {"id":"bl_biscuits",   "label":"Biscuits Station"},
    ]},
    {"id":"fl","label":"FRONTLINE",     "color":AZUL, "sub":"HOSPITALITY","cols":[
        {"id":"fl_comedor",      "label":"Dining Room Station"},
        {"id":"fl_cajero",       "label":"Cashier Station"},
        {"id":"fl_autoservicio", "label":"Drive-Thru Station"},
        {"id":"fl_drive_speed",  "label":"Drive-Thru Speed Team Certification"},
    ]},
    {"id":"ct","label":"CREW TRAINER",  "color":VERDE,"sub":"","cols":[
        {"id":"ct_backline",  "label":"Crew Trainer – BACKLINE"},
        {"id":"ct_frontline", "label":"Crew Trainer – FRONTLINE"},
        {"id":"ct_pic",       "label":"Centerpost Position\n(Certified to cover PIC)"},
    ]},
]

# ── HELPERS ESTILOS ───────────────────────────────────────────
def fill(h):
    return PatternFill("solid", fgColor=h)

def fnt(color=BLANC, bold=False, size=9):
    return Font(name="Arial", size=size, bold=bold, color=color)

def aln(h="center", v="center", wrap=False, rot=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, text_rotation=rot)

def bdr(c="2A2A2A", thick=False):
    s = Side(style="medium" if thick else "thin", color="FFDE21" if thick else c)
    return Border(left=s, right=s, top=s, bottom=s)

def sc(ws, r, c, v, bg, fg=BLANC, bold=False, ha="center", sz=9,
       wrap=False, rot=0, thick=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.fill = fill(bg)
    cell.font = fnt(fg, bold, sz)
    cell.alignment = aln(ha, "center", wrap, rot)
    cell.border = bdr(thick=thick)
    return cell

# ── BUILDER DE HOJA ───────────────────────────────────────────
def build_sheet(wb, grupos, sucursal, entrenador, empleados, titulo):
    all_cols = [c for g in grupos for c in g["cols"]]
    ws = wb.create_sheet(titulo)
    ws.sheet_view.showGridLines = False
    FIXED = 3
    TOTAL = FIXED + len(all_cols) + 1
    HOY   = datetime.date.today().strftime("%d/%m/%Y")

    # ── Fila 1: banner ──────────────────────────────────────
    sc(ws,1,1, f"{titulo}  —  {sucursal.upper()}", NEG, AMAR, True, "left", 13)
    for c in range(2, TOTAL): sc(ws,1,c, "", NEG)
    sc(ws,1,TOTAL, f"Generado: {HOY}", NEG, GRIS, False, "right", 8)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=TOTAL-1)
    ws.row_dimensions[1].height = 26

    # ── Fila 2: sub-info ────────────────────────────────────
    sc(ws,2,1,f"Entrenador: {entrenador}   |   Total empleados: {len(empleados)}",
       NEG3, GRIS, False, "left", 8)
    for c in range(2,TOTAL+1): sc(ws,2,c,"",NEG3)
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=TOTAL)
    ws.row_dimensions[2].height = 14

    # ── Fila 3: separador ───────────────────────────────────
    for c in range(1,TOTAL+1): sc(ws,3,c,"",NEG)
    ws.row_dimensions[3].height = 5

    # ── Fila 4: grupos ──────────────────────────────────────
    for c in range(1,FIXED+1): sc(ws,4,c,"",NEG2)
    ws.merge_cells(start_row=4,start_column=1,end_row=4,end_column=FIXED)
    ci = FIXED+1
    for g in grupos:
        n   = len(g["cols"])
        lbl = g["label"] + ("\n"+g["sub"] if g.get("sub") else "")
        sc(ws,4,ci, lbl, g["color"], BLANC, True, "center", 8, wrap=True)
        for i in range(1,n): sc(ws,4,ci+i,"",g["color"])
        if n>1: ws.merge_cells(start_row=4,start_column=ci,end_row=4,end_column=ci+n-1)
        ci += n
    sc(ws,4,ci,"",NEG2)
    ws.row_dimensions[4].height = 18

    # ── Fila 5: headers rotados 90° ─────────────────────────
    sc(ws,5,1,"FECHA\nINGRESO",        NEG2, AMAR, True, "center", 8, wrap=True)
    sc(ws,5,2,"NOMBRE DEL EMPLEADO",   NEG2, AMAR, True, "center", 9, wrap=True)
    sc(ws,5,3,"ENTRENADOR",            NEG2, AMAR, True, "center", 8, wrap=True)
    ci = FIXED+1
    for g in grupos:
        for col in g["cols"]:
            cell = ws.cell(row=5, column=ci, value=col["label"])
            cell.fill = fill(g["color"])
            cell.font = Font(name="Arial", size=7, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="bottom",
                                       wrap_text=False, text_rotation=90)
            cell.border = bdr()
            ci += 1
    sc(ws,5,ci,"%",NEG2,AMAR,True,"center",9,wrap=True)
    ws.row_dimensions[5].height = 90

    # ── Anchos ──────────────────────────────────────────────
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 19
    for i in range(len(all_cols)):
        ws.column_dimensions[get_column_letter(FIXED+1+i)].width = 4.8
    ws.column_dimensions[get_column_letter(TOTAL)].width = 9

    # ── Filas empleados ─────────────────────────────────────
    R0 = 6
    for i, emp in enumerate(empleados):
        row  = R0+i
        bg   = NEG2 if i%2==0 else NEG3
        star = emp.get("star", {})
        done = sum(1 for c in all_cols if star.get(c["id"]))
        pct  = round(done/len(all_cols)*100) if all_cols else 0
        pFg  = "4CAF50" if pct>=80 else AMAR if pct>=50 else ROJO

        sc(ws,row,1, emp.get("ingreso","—"),   bg, GRIS,  False, "center", 8)
        sc(ws,row,2, emp.get("nombre","—"),    bg, BLANC, True,  "left",   9)
        sc(ws,row,3, emp.get("entrenador","—"),bg, GRIS,  False, "left",   8)

        ci = FIXED+1
        for g in grupos:
            for col in g["cols"]:
                ok  = bool(star.get(col["id"]))
                sbg = "1C1A00" if ok else bg
                sfg = AMAR     if ok else GRIS2
                cell = ws.cell(row=row, column=ci, value="★" if ok else "☆")
                cell.fill = fill(sbg)
                cell.font = Font(name="Arial", size=12, bold=True, color=sfg)
                cell.alignment = aln("center")
                cell.border = bdr()
                ci += 1

        sc(ws,row,ci, f"{pct}%", bg, pFg, True, "center", 9)
        ws.row_dimensions[row].height = 16

    # ── Fila promedio ────────────────────────────────────────
    TR = R0+len(empleados)
    for c in range(1,TOTAL+1): sc(ws,TR,c,"",NEG2,thick=True)
    cell = ws.cell(row=TR,column=2); cell.value="PROMEDIO SUCURSAL"
    cell.font=fnt(AMAR,True,9); cell.alignment=aln("left"); cell.border=bdr(thick=True)
    ws.merge_cells(start_row=TR,start_column=1,end_row=TR,end_column=3)
    ci = FIXED+1; spct=0
    for col in all_cols:
        cnt = sum(1 for e in empleados if e.get("star",{}).get(col["id"]))
        p   = round(cnt/len(empleados)*100) if empleados else 0
        spct += p
        cell = ws.cell(row=TR,column=ci,value=f"{cnt}/{len(empleados)}")
        cell.fill=fill(NEG2); cell.font=Font(name="Arial",size=7,bold=True,color="FFFFFF")
        cell.alignment=aln("center"); cell.border=bdr(thick=True); ci+=1
    avg = round(spct/len(all_cols)) if all_cols else 0
    aFg = "4CAF50" if avg>=80 else AMAR if avg>=50 else ROJO
    sc(ws,TR,ci,f"{avg}%",NEG2,aFg,True,"center",10,thick=True)
    ws.row_dimensions[TR].height = 22

    # ── Pie ─────────────────────────────────────────────────
    for c in range(1,TOTAL+1): sc(ws,TR+1,c,"",NEG)
    ws.row_dimensions[TR+1].height = 5
    sc(ws,TR+2,1,
       f"Cero Tenedor  •  Sistema de Capacitación Carl's Jr.  •  {sucursal.upper()}",
       NEG,GRIS2,False,"left",7)
    for c in range(2,TOTAL+1): sc(ws,TR+2,c,"",NEG)
    ws.merge_cells(start_row=TR+2,start_column=1,end_row=TR+2,end_column=TOTAL)
    ws.row_dimensions[TR+2].height = 12

    ws.freeze_panes = f"{get_column_letter(FIXED+1)}6"

# ── HANDLER VERCEL ────────────────────────────────────────────
class handler(BaseHTTPRequestHandler):

    def do_OPTIONS(self):
        self.send_response(200)
        self._cors()
        self.end_headers()

    def do_POST(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            body   = json.loads(self.rfile.read(length))
        except Exception:
            self.send_response(400)
            self._cors()
            self.send_header("Content-Type","application/json")
            self.end_headers()
            self.wfile.write(json.dumps({"error":"JSON inválido"}).encode())
            return

        tipo       = body.get("tipo", "star")          # "star" | "matriz"
        sucursal   = body.get("sucursal","Sucursal")
        entrenador = body.get("entrenador","Entrenador")
        empleados  = body.get("empleados", [])

        grupos = SP_GRUPOS if tipo == "star" else MZ_GRUPOS
        titulo = "STAR PERFORMANCE" if tipo == "star" else "MATRIZ DE CAPACITACIÓN"
        nombre_archivo = (
            f"StarPerformance_{sucursal.replace(' ','_')}_{datetime.date.today()}.xlsx"
            if tipo == "star"
            else f"Matriz_{sucursal.replace(' ','_')}_{datetime.date.today()}.xlsx"
        )

        wb = Workbook()
        wb.remove(wb.active)
        build_sheet(wb, grupos, sucursal, entrenador, empleados, titulo)

        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        self.send_response(200)
        self._cors()
        self.send_header("Content-Type",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition",
            f'attachment; filename="{nombre_archivo}"')
        self.send_header("Content-Length", str(len(xlsx_bytes)))
        self.end_headers()
        self.wfile.write(xlsx_bytes)

    def _cors(self):
        self.send_header("Access-Control-Allow-Origin",  "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
